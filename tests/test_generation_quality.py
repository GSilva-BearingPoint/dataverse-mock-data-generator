from __future__ import annotations

import copy
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from build_testdata_workbook import build_workbook
import check_dependencies
from discover_testdata_workbook import discover_plan, inspect_table, parse_table_counts
from progress import progress_message
from quality_profile import build_quality_profile
from summarize_testdata_workbook import summarize
from testdata_common import generate_records
import validate_testdata_workbook
from validate_testdata_workbook import validate


def column(
    name: str,
    attr_type: str = "String",
    required: str = "None",
    *,
    custom: bool = True,
    fmt: str | None = None,
    max_length: int | None = 100,
) -> dict:
    value = {
        "LogicalName": name,
        "AttributeType": attr_type,
        "RequiredLevel": {"Value": required, "CanBeChanged": False},
        "DisplayName": {"UserLocalizedLabel": {"Label": name}},
        "IsValidForCreate": True,
        "IsValidODataAttribute": True,
        "IsCustomAttribute": custom,
        "AttributeOf": None,
    }
    if fmt:
        value["Format"] = fmt
    if attr_type in {"String", "Memo"} and max_length is not None:
        value["MaxLength"] = max_length
    return value


class FakeResponse:
    def __init__(self, values: list[dict]):
        self.values = values

    def json(self) -> dict:
        return {"value": self.values}


class FakeOData:
    api = "https://metadata.invalid/api/data/v9.2"

    def __init__(self, schemas: dict[str, dict]):
        self.schemas = schemas

    def _request(self, _method: str, url: str) -> FakeResponse:
        table_name = url.split("LogicalName='", 1)[1].split("'", 1)[0]
        flags = [
            {
                key: raw[key]
                for key in (
                    "LogicalName",
                    "IsValidForCreate",
                    "IsValidODataAttribute",
                    "AttributeOf",
                    "IsCustomAttribute",
                )
            }
            for raw in self.schemas[table_name]["columns"]
        ]
        return FakeResponse(flags)


class FakeTables:
    def __init__(self, schemas: dict[str, dict]):
        self.schemas = schemas

    def get(self, table_name: str) -> dict:
        return copy.deepcopy(self.schemas[table_name]["table"])

    def list_columns(self, table_name: str, typed: bool = True) -> list[dict]:
        self.assert_true(typed)
        return copy.deepcopy(self.schemas[table_name]["columns"])

    def get_alternate_keys(self, table_name: str) -> list[dict]:
        return copy.deepcopy(self.schemas[table_name].get("keys", []))

    def list_table_relationships(self, table_name: str) -> list[dict]:
        return copy.deepcopy(self.schemas[table_name].get("relationships", []))

    @staticmethod
    def assert_true(value: bool) -> None:
        if not value:
            raise AssertionError("typed metadata was not requested")


class FakeRecords:
    def __init__(self, existing: dict[str, list[dict]] | None = None):
        self.existing = existing or {}

    def list(self, table_name: str, select: list[str], top: int) -> list[dict]:
        return copy.deepcopy(self.existing.get(table_name, []))[:top]


class FakeClient:
    def __init__(
        self,
        schemas: dict[str, dict],
        existing: dict[str, list[dict]] | None = None,
    ):
        self.schemas = schemas
        self.tables = FakeTables(schemas)
        self.records = FakeRecords(existing)
        self.odata = FakeOData(schemas)

    def _get_odata(self) -> FakeOData:
        return self.odata


def table_schema(
    name: str,
    columns: list[dict],
    *,
    primary_name: str,
    relationships: list[dict] | None = None,
    keys: list[dict] | None = None,
) -> dict:
    return {
        "table": {
            "LogicalName": name,
            "EntitySetName": f"{name}s",
            "PrimaryIdAttribute": f"{name}id",
            "PrimaryNameAttribute": primary_name,
            "DisplayName": {"UserLocalizedLabel": {"Label": name}},
        },
        "columns": columns,
        "relationships": relationships or [],
        "keys": keys or [],
    }


class DiscoveryTests(unittest.TestCase):
    def test_fixed_required_primary_and_representative_custom_columns(self) -> None:
        schemas = {
            "new_widget": table_schema(
                "new_widget",
                [
                    column("new_name", required="SystemRequired"),
                    column("new_code"),
                    column("new_category"),
                    column("new_description", attr_type="Memo"),
                    column("new_extra"),
                ],
                primary_name="new_name",
            )
        }
        info = inspect_table(FakeClient(schemas), "new_widget", 2, False)
        selected = {item["logical_name"]: item for item in info["columns"]}

        self.assertTrue(selected["new_name"]["required"])
        self.assertIn("primary_name", selected["new_name"]["selection_reasons"])
        representative = [
            item
            for item in info["columns"]
            if "representative_custom" in item["selection_reasons"]
        ]
        self.assertEqual(3, len(representative))
        self.assertEqual(4, len(info["columns"]))

    def test_required_dependency_is_filled_and_explained(self) -> None:
        schemas = {
            "new_child": table_schema(
                "new_child",
                [
                    column("new_name", required="ApplicationRequired"),
                    column("new_parentid", attr_type="Lookup", required="ApplicationRequired"),
                ],
                primary_name="new_name",
                relationships=[
                    {
                        "ReferencingEntity": "new_child",
                        "ReferencingAttribute": "new_parentid",
                        "ReferencedEntity": "new_parent",
                        "ReferencingEntityNavigationPropertyName": "new_parentid_nav",
                    }
                ],
            ),
            "new_parent": table_schema(
                "new_parent",
                [
                    column("new_name", required="ApplicationRequired"),
                    column("new_phone", fmt="Phone"),
                ],
                primary_name="new_name",
            ),
        }
        plan = discover_plan(
            FakeClient(schemas),
            ["new_child"],
            2,
            None,
            "https://environment-a.invalid",
            table_counts={"new_child": 4},
        )
        by_name = {table["logical_name"]: table for table in plan["tables"]}

        self.assertEqual(4, by_name["new_child"]["count"])
        self.assertEqual(2, by_name["new_parent"]["count"])
        self.assertEqual("implicit", by_name["new_parent"]["inclusion"]["kind"])
        reason = by_name["new_parent"]["inclusion"]["reasons"][0]
        self.assertEqual("new_child", reason["source_table"])
        self.assertEqual("new_parentid", reason["lookup_column"])
        self.assertIn("new_parentid", by_name["new_child"]["quality_profile"]["core_columns"])

    def test_profile_decisions_do_not_depend_on_environment_identity(self) -> None:
        schemas = {
            "abc_item": table_schema(
                "abc_item",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_phone", fmt="Phone"),
                    column("abc_notes", attr_type="Memo"),
                ],
                primary_name="abc_name",
            )
        }
        first = discover_plan(
            FakeClient(schemas),
            ["abc_item"],
            3,
            None,
            "https://environment-a.invalid",
        )
        second = discover_plan(
            FakeClient(schemas),
            ["abc_item"],
            3,
            None,
            "https://environment-b.invalid",
        )

        self.assertEqual(first["tables"], second["tables"])

    def test_table_count_parser_rejects_invalid_values(self) -> None:
        self.assertEqual({"a": 2, "b": 5}, parse_table_counts("a=2,b=5"))
        for value in ("a", "a=0", "a=x", "a=1,a=2"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                parse_table_counts(value)

    def test_confirmed_relationship_rule_selects_optional_dependency(self) -> None:
        schemas = {
            "abc_order": table_schema(
                "abc_order",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_customerid", attr_type="Lookup"),
                ],
                primary_name="abc_name",
                relationships=[
                    {
                        "ReferencingEntity": "abc_order",
                        "ReferencingAttribute": "abc_customerid",
                        "ReferencedEntity": "abc_customer",
                        "ReferencingEntityNavigationPropertyName": "abc_customerid_nav",
                    }
                ],
            ),
            "abc_customer": table_schema(
                "abc_customer",
                [column("abc_name", required="ApplicationRequired")],
                primary_name="abc_name",
            ),
        }
        plan = discover_plan(
            FakeClient(schemas),
            ["abc_order"],
            2,
            None,
            "https://environment.invalid",
            context={
                "tables": {
                    "abc_order": {
                        "relationships": {"abc_customerid": "abc_customer"}
                    }
                }
            },
        )
        by_name = {table["logical_name"]: table for table in plan["tables"]}

        self.assertEqual("implicit", by_name["abc_customer"]["inclusion"]["kind"])
        self.assertEqual(
            "selected_rule",
            by_name["abc_customer"]["inclusion"]["reasons"][0]["requirement"],
        )
        self.assertEqual(
            "abc_customer",
            by_name["abc_order"]["lookups"][0]["target_table"],
        )

    def test_existing_record_dependency_is_not_generated(self) -> None:
        schemas = {
            "abc_order": table_schema(
                "abc_order",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_customerid", attr_type="Lookup", required="ApplicationRequired"),
                ],
                primary_name="abc_name",
                relationships=[
                    {
                        "ReferencingEntity": "abc_order",
                        "ReferencingAttribute": "abc_customerid",
                        "ReferencedEntity": "abc_customer",
                        "ReferencingEntityNavigationPropertyName": "abc_customerid_nav",
                    }
                ],
            ),
            "abc_customer": table_schema(
                "abc_customer",
                [column("abc_name", required="ApplicationRequired")],
                primary_name="abc_name",
            ),
        }
        client = FakeClient(
            schemas,
            existing={"abc_customer": [{"abc_customerid": "00000000-0000-0000-0000-000000000001"}]},
        )
        plan = discover_plan(
            client,
            ["abc_order"],
            2,
            None,
            "https://environment.invalid",
            existing_tables={"abc_customer"},
        )

        self.assertEqual(["abc_order"], [table["logical_name"] for table in plan["tables"]])
        self.assertEqual(
            ["00000000-0000-0000-0000-000000000001"],
            plan["tables"][0]["lookups"][0]["existing_ids"],
        )

    def test_required_polymorphic_lookup_requires_an_explicit_target(self) -> None:
        schemas = {
            "abc_activity": table_schema(
                "abc_activity",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_regardingid", attr_type="Lookup", required="ApplicationRequired"),
                ],
                primary_name="abc_name",
                relationships=[
                    {
                        "ReferencingEntity": "abc_activity",
                        "ReferencingAttribute": "abc_regardingid",
                        "ReferencedEntity": target,
                        "ReferencingEntityNavigationPropertyName": f"abc_regardingid_{target}",
                    }
                    for target in ("abc_account", "abc_case")
                ],
            ),
            "abc_account": table_schema(
                "abc_account",
                [column("abc_name", required="ApplicationRequired")],
                primary_name="abc_name",
            ),
            "abc_case": table_schema(
                "abc_case",
                [column("abc_name", required="ApplicationRequired")],
                primary_name="abc_name",
            ),
        }
        with self.assertRaisesRegex(ValueError, "required lookup"):
            discover_plan(
                FakeClient(schemas),
                ["abc_activity"],
                2,
                None,
                "https://environment.invalid",
            )

    def test_all_writable_mode_includes_every_supported_custom_column(self) -> None:
        schemas = {
            "abc_item": table_schema(
                "abc_item",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_one"),
                    column("abc_two"),
                    column("abc_three"),
                    column("abc_four"),
                ],
                primary_name="abc_name",
            )
        }
        info = inspect_table(FakeClient(schemas), "abc_item", 2, True)

        self.assertEqual(5, len(info["columns"]))
        self.assertTrue(
            all("all_writable" in item["selection_reasons"] for item in info["columns"])
        )

    def test_cycle_removal_preserves_optional_edges_outside_the_cycle(self) -> None:
        schemas = {
            "abc_a": table_schema(
                "abc_a",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_bid", attr_type="Lookup", required="ApplicationRequired"),
                ],
                primary_name="abc_name",
                relationships=[
                    {
                        "ReferencingEntity": "abc_a",
                        "ReferencingAttribute": "abc_bid",
                        "ReferencedEntity": "abc_b",
                        "ReferencingEntityNavigationPropertyName": "abc_bid_nav",
                    }
                ],
            ),
            "abc_b": table_schema(
                "abc_b",
                [
                    column("abc_name", required="ApplicationRequired"),
                    column("abc_aid", attr_type="Lookup"),
                    column("abc_cid", attr_type="Lookup"),
                ],
                primary_name="abc_name",
                relationships=[
                    {
                        "ReferencingEntity": "abc_b",
                        "ReferencingAttribute": "abc_aid",
                        "ReferencedEntity": "abc_a",
                        "ReferencingEntityNavigationPropertyName": "abc_aid_nav",
                    },
                    {
                        "ReferencingEntity": "abc_b",
                        "ReferencingAttribute": "abc_cid",
                        "ReferencedEntity": "abc_c",
                        "ReferencingEntityNavigationPropertyName": "abc_cid_nav",
                    },
                ],
            ),
            "abc_c": table_schema(
                "abc_c",
                [column("abc_name", required="ApplicationRequired")],
                primary_name="abc_name",
            ),
        }
        plan = discover_plan(
            FakeClient(schemas),
            ["abc_a", "abc_c"],
            2,
            None,
            "https://environment.invalid",
        )
        by_name = {table["logical_name"]: table for table in plan["tables"]}
        lookup_columns = {lookup["column"] for lookup in by_name["abc_b"]["lookups"]}

        self.assertNotIn("abc_aid", lookup_columns)
        self.assertIn("abc_cid", lookup_columns)


class WorkbookTests(unittest.TestCase):
    def _spec(self) -> dict:
        table = {
            "logical_name": "abc_widget",
            "entity_set_name": "abc_widgets",
            "primary_id_column": "abc_widgetid",
            "primary_name_column": "abc_name",
            "count": 2,
            "columns": [
                {
                    "logical_name": "abc_name",
                    "label": "Name",
                    "type": "String",
                    "required": True,
                    "unique": True,
                    "max_length": 100,
                    "selection_reasons": ["required", "primary_name"],
                },
                {
                    "logical_name": "abc_phone",
                    "label": "Phone",
                    "type": "String",
                    "role": "phone",
                    "max_length": 100,
                    "selection_reasons": ["semantic_profile"],
                },
                {
                    "logical_name": "abc_description",
                    "label": "Description",
                    "type": "Memo",
                    "role": "description",
                    "max_length": 500,
                    "selection_reasons": ["semantic_profile"],
                },
            ],
            "lookups": [],
            "context_rules": {},
            "inclusion": {
                "kind": "implicit",
                "dependency_path": ["abc_parent", "abc_widget"],
                "reasons": [
                    {
                        "source_table": "abc_parent",
                        "lookup_column": "abc_widgetid",
                        "requirement": "required",
                        "target_table": "abc_widget",
                    }
                ],
            },
        }
        table["quality_profile"] = build_quality_profile(table)
        return {
            "name": "test",
            "environment": "https://environment.invalid",
            "run_id": "unit-test",
            "warnings": ["Synthetic warning"],
            "tables": [table],
        }

    def test_validator_rejects_bare_bones_rows_and_duplicate_keys(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "quality.xlsx"
            build_workbook(self._spec(), path)
            _rows, initial_errors = validate(path)
            self.assertEqual([], initial_errors)

            workbook = load_workbook(path)
            worksheet = workbook["abc_widget"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(2, headers.index("abc_phone") + 1).value = None
            worksheet.cell(2, headers.index("abc_description") + 1).value = None
            worksheet.cell(3, headers.index("abc_name") + 1).value = worksheet.cell(
                2, headers.index("abc_name") + 1
            ).value
            workbook.save(path)

            _rows, errors = validate(path)
            self.assertTrue(
                any("record is not meaningful enough" in error for error in errors),
                errors,
            )
            self.assertTrue(
                any("alternate key (abc_name) has duplicate values" in error for error in errors),
                errors,
            )

    def test_summary_contains_samples_completeness_and_dependency_reason(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "summary.xlsx"
            build_workbook(self._spec(), path)
            output = summarize(path)

            self.assertIn("[Step 7/8 - Approval] WAITING:", output)
            self.assertIn("Table abc_widget: 2 record(s), implicit", output)
            self.assertIn("abc_parent.abc_widgetid required abc_widget", output)
            self.assertIn("path: abc_parent -> abc_widget", output)
            self.assertIn("Completeness:", output)
            self.assertIn("Sample 1:", output)
            self.assertIn("Synthetic warning", output)

    def test_validator_checks_choice_length_and_numeric_bounds(self) -> None:
        spec = self._spec()
        table = spec["tables"][0]
        table["columns"].extend(
            [
                {
                    "logical_name": "abc_tier",
                    "label": "Tier",
                    "type": "Picklist",
                    "options": [{"value": 1, "label": "One"}, {"value": 2, "label": "Two"}],
                    "selection_reasons": ["recommended"],
                },
                {
                    "logical_name": "abc_score",
                    "label": "Score",
                    "type": "Integer",
                    "min_value": 1,
                    "max_value": 10,
                    "selection_reasons": ["representative_custom"],
                },
                {
                    "logical_name": "abc_active",
                    "label": "Active",
                    "type": "Boolean",
                    "selection_reasons": ["representative_custom"],
                },
                {
                    "logical_name": "abc_date",
                    "label": "Date",
                    "type": "DateTime",
                    "format": "DateOnly",
                    "selection_reasons": ["representative_custom"],
                },
            ]
        )
        table["quality_profile"] = build_quality_profile(table)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "constraints.xlsx"
            build_workbook(spec, path)
            workbook = load_workbook(path)
            worksheet = workbook["abc_widget"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(2, headers.index("abc_tier") + 1).value = 99
            worksheet.cell(2, headers.index("abc_score") + 1).value = 11
            worksheet.cell(2, headers.index("abc_phone") + 1).value = "x" * 101
            worksheet.cell(2, headers.index("abc_active") + 1).value = "yes"
            worksheet.cell(2, headers.index("abc_date") + 1).value = "not-a-date"
            workbook.save(path)

            _rows, errors = validate(path)
            self.assertTrue(any("invalid choice value" in error for error in errors), errors)
            self.assertTrue(any("exceeds maximum 10" in error for error in errors), errors)
            self.assertTrue(any("exceeds maximum length 100" in error for error in errors), errors)
            self.assertTrue(any("must be a Boolean value" in error for error in errors), errors)
            self.assertTrue(any("must be a valid date" in error for error in errors), errors)

    def test_required_secondary_address_lines_are_populated(self) -> None:
        spec = self._spec()
        table = spec["tables"][0]
        table["columns"].extend(
            [
                {
                    "logical_name": "abc_address1_line2",
                    "label": "Address line 2",
                    "type": "String",
                    "role": "addr_line2",
                    "required": True,
                    "selection_reasons": ["required"],
                },
                {
                    "logical_name": "abc_address1_line3",
                    "label": "Address line 3",
                    "type": "String",
                    "role": "addr_line3",
                    "required": True,
                    "selection_reasons": ["required"],
                },
            ]
        )
        table["quality_profile"] = build_quality_profile(table)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "required-address.xlsx"
            build_workbook(spec, path)
            rows, errors = validate(path)

            self.assertEqual([], errors)
            self.assertTrue(rows["abc_widget"][0]["abc_address1_line2"])
            self.assertTrue(rows["abc_widget"][0]["abc_address1_line3"])

    def test_composite_alternate_keys_are_validated_as_tuples(self) -> None:
        table = {
            "logical_name": "abc_regioncode",
            "entity_set_name": "abc_regioncodes",
            "primary_id_column": "abc_regioncodeid",
            "primary_name_column": "abc_name",
            "count": 2,
            "columns": [
                {
                    "logical_name": "abc_name",
                    "label": "Name",
                    "type": "String",
                    "required": True,
                    "selection_reasons": ["required", "primary_name"],
                },
                {
                    "logical_name": "abc_country",
                    "label": "Country",
                    "type": "String",
                    "selection_reasons": ["representative_custom"],
                },
                {
                    "logical_name": "abc_code",
                    "label": "Code",
                    "type": "String",
                    "selection_reasons": ["representative_custom"],
                },
            ],
            "alternate_keys": [["abc_country", "abc_code"]],
            "lookups": [],
            "context_rules": {},
            "inclusion": {"kind": "requested", "dependency_path": ["abc_regioncode"], "reasons": []},
            "records": [
                {"abc_name": "One", "abc_country": "PT", "abc_code": "1"},
                {"abc_name": "Two", "abc_country": "PT", "abc_code": "2"},
            ],
        }
        table["quality_profile"] = build_quality_profile(table)
        spec = {
            "name": "test",
            "environment": "https://environment.invalid",
            "run_id": "composite-key",
            "tables": [table],
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "composite.xlsx"
            build_workbook(spec, path)
            _rows, errors = validate(path)
            self.assertEqual([], errors)

            workbook = load_workbook(path)
            worksheet = workbook["abc_regioncode"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(3, headers.index("abc_code") + 1).value = "1"
            workbook.save(path)
            _rows, errors = validate(path)
            self.assertTrue(
                any("alternate key (abc_country, abc_code)" in error for error in errors),
                errors,
            )

    def test_generated_composite_key_duplicates_are_adjusted(self) -> None:
        table = {
            "logical_name": "abc_external",
            "count": 3,
            "alternate_keys": [["abc_country", "abc_code"]],
            "columns": [
                {
                    "logical_name": "abc_country",
                    "type": "String",
                    "max_length": 10,
                },
                {
                    "logical_name": "abc_code",
                    "type": "String",
                    "max_length": 10,
                },
            ],
            "context_rules": {
                "defaults": {"abc_country": "PT", "abc_code": "EXT"}
            },
        }

        records = generate_records(table, None, "composite-generation")
        keys = {(row["abc_country"], row["abc_code"]) for row in records}

        self.assertEqual(3, len(keys))
        self.assertEqual("EXT", records[0]["abc_code"])
        self.assertEqual("EXT-2", records[1]["abc_code"])


class ProgressVisibilityTests(unittest.TestCase):
    def test_progress_message_explains_current_and_next_step(self) -> None:
        message = progress_message(
            5,
            "running",
            "Generating metadata-driven records.",
            "Validate the workbook.",
        )

        self.assertEqual(
            "[Step 5/8 - Record generation] RUNNING: Generating metadata-driven records. "
            "Next: Validate the workbook.",
            message,
        )

    def test_dependency_json_remains_machine_readable_with_progress(self) -> None:
        ready = {
            "python": "python",
            "missing_tools": [],
            "missing_python_modules": [],
            "ready": True,
        }
        output = io.StringIO()
        with (
            patch.object(check_dependencies, "check", return_value=ready),
            patch.object(sys, "argv", ["check_dependencies.py", "--json"]),
            redirect_stdout(output),
            self.assertRaises(SystemExit) as exit_context,
        ):
            check_dependencies.main()

        result = json.loads(output.getvalue())
        self.assertEqual(0, exit_context.exception.code)
        self.assertEqual(1, result["progress"]["step"])
        self.assertEqual("complete", result["progress"]["status"])

    def test_dependency_json_reports_blocked_state(self) -> None:
        blocked = {
            "python": "python",
            "missing_tools": [{"name": "pac", "install": "install"}],
            "missing_python_modules": [],
            "ready": False,
        }
        output = io.StringIO()
        with (
            patch.object(check_dependencies, "check", return_value=blocked),
            patch.object(sys, "argv", ["check_dependencies.py", "--json"]),
            redirect_stdout(output),
            self.assertRaises(SystemExit) as exit_context,
        ):
            check_dependencies.main()

        result = json.loads(output.getvalue())
        self.assertEqual(1, exit_context.exception.code)
        self.assertEqual("blocked", result["progress"]["status"])

    def test_validation_cli_reports_running_and_complete_states(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "progress.xlsx"
            build_workbook(WorkbookTests()._spec(), path)
            output = io.StringIO()
            with (
                patch.object(
                    sys,
                    "argv",
                    ["validate_testdata_workbook.py", "--workbook", str(path)],
                ),
                redirect_stdout(output),
            ):
                validate_testdata_workbook.main()

            text = output.getvalue()
            self.assertIn("[Step 6/8 - Workbook validation] RUNNING:", text)
            self.assertIn("[Step 6/8 - Workbook validation] COMPLETE:", text)


if __name__ == "__main__":
    unittest.main()
