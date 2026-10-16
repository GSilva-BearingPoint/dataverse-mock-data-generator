from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from derive_values import derive_values
from quality_profile import evaluate_record
from testdata_common import (
    INTERNAL_KEY,
    dependency_order,
    generate_records,
    load_json,
    run_id,
    safe_sheet_name,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
LOOKUP_FILL = PatternFill("solid", fgColor="FFF2CC")
REQUIRED_FILL = PatternFill("solid", fgColor="FCE4D6")
METADATA_CHUNK_SIZE = 30000


def _append_metadata(metadata, kind: str, name: str, value: dict[str, Any]) -> None:
    raw = json.dumps(value)
    for part, start in enumerate(range(0, len(raw), METADATA_CHUNK_SIZE)):
        metadata.append([kind, name, part, raw[start:start + METADATA_CHUNK_SIZE]])


def _table_records(
    table: dict[str, Any], seed_email: str | None, current_run_id: str
) -> list[dict[str, Any]]:
    records = table.get("records")
    if records is None:
        return generate_records(table, seed_email, current_run_id)
    for index, record in enumerate(records):
        record.setdefault(INTERNAL_KEY, f"TD-{table['logical_name']}-{index + 1:04d}")
    return records


def _inclusion_reason(table: dict[str, Any]) -> str:
    inclusion = table.get("inclusion") or {}
    reasons = inclusion.get("reasons") or []
    if not reasons:
        return "Explicitly requested"
    reason_text = "; ".join(
        f"{reason['source_table']}.{reason['lookup_column']} "
        f"({reason['requirement']})"
        for reason in reasons
    )
    path = inclusion.get("dependency_path") or []
    return reason_text + (f"; path: {' -> '.join(path)}" if path else "")


def build_workbook(spec: dict[str, Any], output_path: Path) -> None:
    tables = spec["tables"]
    current_run_id = run_id(spec)
    seed_email = spec.get("seed_email")
    order = dependency_order(tables)
    by_name = {table["logical_name"]: table for table in tables}
    records_by_table = {
        name: _table_records(by_name[name], seed_email, current_run_id) for name in order
    }
    for name in order:
        for lookup in by_name[name].get("lookups", []):
            target_records = records_by_table.get(lookup["target_table"])
            pool = lookup.get("existing_ids")
            if not pool and not target_records:
                if lookup.get("required"):
                    raise ValueError(
                        f"Required lookup {name}.{lookup['column']} targets {lookup['target_table']}, "
                        "which has no generated records."
                    )
                continue
            for index, record in enumerate(records_by_table[name]):
                record[lookup["column"]] = (
                    pool[index % len(pool)] if pool
                    else target_records[index % len(target_records)][INTERNAL_KEY]
                )
    derive_values(by_name, records_by_table, order)
    for name in order:
        table = by_name[name]
        for index, record in enumerate(records_by_table[name], start=1):
            quality_errors = evaluate_record(table, record)
            if quality_errors:
                raise ValueError(
                    f"{name} generated row {index} does not satisfy its quality profile: "
                    + "; ".join(quality_errors)
                )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "README"
    summary.append(["Dataverse Test Data Plan"])
    summary.append(["Environment", spec.get("environment", "")])
    summary.append(["Run ID", current_run_id])
    summary.append(["Seed email", seed_email or "(not required)"])
    summary.append(["Repository context", spec.get("context_source") or "(none provided)"])
    summary.append(["Approval status", "NOT APPROVED - workbook is a preview only"])
    summary.append([])
    summary.append(
        ["Creation order", "Table", "Suggested records", "Worksheet", "Source", "Reason"]
    )

    used_sheet_names = {"readme", "_metadata"}
    sheet_by_table = {
        name: safe_sheet_name(by_name[name].get("sheet_name") or name, used_sheet_names)
        for name in order
    }
    for level, name in enumerate(order, start=1):
        table = by_name[name]
        inclusion = table.get("inclusion") or {}
        summary.append(
            [
                level,
                name,
                len(records_by_table[name]),
                sheet_by_table[name],
                inclusion.get("kind", "requested"),
                _inclusion_reason(table),
            ]
        )
    implicit_tables = [
        by_name[name]
        for name in order
        if (by_name[name].get("inclusion") or {}).get("kind") == "implicit"
    ]
    if implicit_tables:
        summary.append([])
        summary.append(["Automatically added tables"])
        for table in implicit_tables:
            summary.append([table["logical_name"], _inclusion_reason(table)])
    warnings = spec.get("warnings") or []
    if warnings:
        summary.append([])
        summary.append(["Warnings requiring review"])
        for warning in warnings:
            summary.append([warning])
    summary.freeze_panes = "A9"
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 42

    metadata = workbook.create_sheet("_Metadata")
    metadata.sheet_state = "veryHidden"
    metadata.append(["kind", "name", "part", "json"])
    _append_metadata(metadata, "run", current_run_id, {
        "name": spec.get("name", "testdata"),
        "run_id": current_run_id,
        "environment": spec.get("environment"),
        "seed_email": seed_email,
        "context_source": spec.get("context_source"),
        "context_notes": spec.get("context_notes") or [],
        "warnings": warnings,
        "order": order,
        "sheet_by_table": sheet_by_table,
    })

    key_cells: dict[str, dict[str, str]] = {}
    for name in order:
        table = by_name[name]
        worksheet = workbook.create_sheet(sheet_by_table[name])
        columns = [INTERNAL_KEY]
        columns.extend(column["logical_name"] for column in table.get("columns", []))
        columns.extend(lookup["column"] for lookup in table.get("lookups", []))
        worksheet.append(columns)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

        column_metadata = {column["logical_name"]: column for column in table.get("columns", [])}
        lookup_metadata = {lookup["column"]: lookup for lookup in table.get("lookups", [])}
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
            if cell.value in lookup_metadata:
                cell.fill = LOOKUP_FILL
                cell.font = Font(color="000000", bold=True)
                target = lookup_metadata[cell.value]["target_table"]
                cell.comment = Comment(f"Lookup to {target}; cells link to the related sheet row.", "Copilot")
            elif column_metadata.get(cell.value, {}).get("required"):
                cell.fill = REQUIRED_FILL
                cell.font = Font(color="000000", bold=True)

        records = records_by_table[name]
        key_cells[name] = {}
        for row_number, record in enumerate(records, start=2):
            worksheet.append([record.get(column) for column in columns])
            key_cells[name][record[INTERNAL_KEY]] = f"A{row_number}"

        for lookup in table.get("lookups", []):
            target_name = lookup["target_table"]
            if lookup.get("existing_ids") or not records_by_table.get(target_name):
                continue
            column_number = columns.index(lookup["column"]) + 1
            for index, record in enumerate(records):
                target_key = record[lookup["column"]]
                cell = worksheet.cell(row=index + 2, column=column_number)
                target_cell = key_cells[target_name][target_key]
                target_sheet = sheet_by_table[target_name].replace("'", "''")
                cell.hyperlink = f"#'{target_sheet}'!{target_cell}"
                cell.style = "Hyperlink"

        for column_number, column_name in enumerate(columns, start=1):
            column = column_metadata.get(column_name)
            if column and column.get("options"):
                values = [
                    option["value"] if isinstance(option, dict) else option
                    for option in column["options"]
                ]
                formula = f'"{",".join(str(value) for value in values)}"'
                if len(formula) <= 255:
                    validation = DataValidation(
                        type="list",
                        formula1=formula,
                        allow_blank=not column.get("required", False),
                    )
                    worksheet.add_data_validation(validation)
                    validation.add(
                        f"{get_column_letter(column_number)}2:"
                        f"{get_column_letter(column_number)}{len(records) + 1}"
                    )
            width = max(14, min(45, len(column_name) + 4))
            worksheet.column_dimensions[get_column_letter(column_number)].width = width

        _append_metadata(metadata, "table", name, {
            **table,
            "sheet_name": sheet_by_table[name],
        })

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build an editable Excel preview for a Dataverse test-data plan."
    )
    parser.add_argument("--plan", required=True, help="Path to the JSON test-data plan.")
    parser.add_argument("--output", required=True, help="Output .xlsx path.")
    args = parser.parse_args()
    build_workbook(load_json(args.plan), Path(args.output))
    print(f"Workbook written: {Path(args.output).resolve()}")


if __name__ == "__main__":
    main()
