from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from progress import emit_progress
from quality_profile import evaluate_record
from testdata_common import INTERNAL_KEY


def load_metadata(workbook) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    metadata = workbook["_Metadata"]
    run: dict[str, Any] = {}
    tables: dict[str, dict[str, Any]] = {}
    chunks: dict[tuple[str, str], list[tuple[int, str]]] = {}
    for kind, name, part, raw_json in metadata.iter_rows(min_row=2, values_only=True):
        chunks.setdefault((kind, name), []).append((int(part), raw_json))
    for (kind, name), parts in chunks.items():
        raw_json = "".join(value for _part, value in sorted(parts))
        value = json.loads(raw_json)
        if kind == "run":
            run = value
        elif kind == "table":
            tables[name] = value
    return run, tables


def read_rows(workbook, table: dict[str, Any]) -> list[dict[str, Any]]:
    worksheet = workbook[table["sheet_name"]]
    headers = [cell.value for cell in worksheet[1]]
    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None and value != "" for value in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def validate(path: Path) -> tuple[dict[str, list[dict[str, Any]]], list[str]]:
    workbook = load_workbook(path, data_only=False)
    _run, tables = load_metadata(workbook)
    rows_by_table = {name: read_rows(workbook, table) for name, table in tables.items()}
    errors: list[str] = []

    keys_by_table: dict[str, set[str]] = {}
    for name, rows in rows_by_table.items():
        keys = [row.get(INTERNAL_KEY) for row in rows]
        missing = [index + 2 for index, key in enumerate(keys) if not key]
        duplicates = sorted({key for key in keys if key and keys.count(key) > 1})
        if missing:
            errors.append(f"{name}: missing {INTERNAL_KEY} on Excel rows {missing}.")
        if duplicates:
            errors.append(f"{name}: duplicate {INTERNAL_KEY} values: {duplicates}.")
        keys_by_table[name] = {str(key) for key in keys if key}

        table = tables[name]
        required = [
            column["logical_name"]
            for column in table.get("columns", [])
            if column.get("required")
        ]
        for row_number, row in enumerate(rows, start=2):
            for column in required:
                if row.get(column) is None or row.get(column) == "":
                    errors.append(f"{name}!{row_number}: required column {column} is blank.")
            for column in table.get("columns", []):
                column_name = column["logical_name"]
                value = row.get(column_name)
                if value is None or value == "":
                    continue
                max_length = column.get("max_length")
                if max_length is not None and len(str(value)) > int(max_length):
                    errors.append(
                        f"{name}!{row_number}: {column_name} exceeds maximum length {max_length}."
                    )
                options = column.get("options") or []
                if options:
                    allowed = {
                        str(option["value"] if isinstance(option, dict) else option)
                        for option in options
                    }
                    if str(value) not in allowed:
                        errors.append(
                            f"{name}!{row_number}: {column_name} has invalid choice value {value!r}."
                        )
                if column.get("type") in {
                    "BigInt",
                    "Decimal",
                    "Double",
                    "Integer",
                    "Money",
                }:
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        errors.append(
                            f"{name}!{row_number}: {column_name} must be numeric."
                        )
                    elif column.get("type") in {"BigInt", "Integer"} and not isinstance(value, int):
                        errors.append(
                            f"{name}!{row_number}: {column_name} must be an integer."
                        )
                    else:
                        minimum = column.get("min_value")
                        maximum = column.get("max_value")
                        if minimum is not None and value < minimum:
                            errors.append(
                                f"{name}!{row_number}: {column_name} is below minimum {minimum}."
                            )
                        if maximum is not None and value > maximum:
                            errors.append(
                                f"{name}!{row_number}: {column_name} exceeds maximum {maximum}."
                            )
                if column.get("type") == "Boolean" and not isinstance(value, bool):
                    errors.append(
                        f"{name}!{row_number}: {column_name} must be a Boolean value."
                    )
                if column.get("type") == "DateTime":
                    valid_date = isinstance(value, (dt.date, dt.datetime))
                    if isinstance(value, str):
                        try:
                            if column.get("format") == "DateOnly":
                                dt.date.fromisoformat(value)
                            else:
                                dt.datetime.fromisoformat(value.replace("Z", "+00:00"))
                            valid_date = True
                        except ValueError:
                            valid_date = False
                    if not valid_date:
                        errors.append(
                            f"{name}!{row_number}: {column_name} must be a valid "
                            + ("date." if column.get("format") == "DateOnly" else "date/time.")
                        )
            for quality_error in evaluate_record(table, row):
                errors.append(
                    f"{name}!{row_number}: record is not meaningful enough: {quality_error}."
                )
        alternate_keys = table.get("alternate_keys") or [
            [column["logical_name"]]
            for column in table.get("columns", [])
            if column.get("unique")
        ]
        for key_columns in alternate_keys:
            values = [
                tuple(str(row[column_name]) for column_name in key_columns)
                for row in rows
                if all(
                    row.get(column_name) is not None and row.get(column_name) != ""
                    for column_name in key_columns
                )
            ]
            duplicates = sorted({value for value in values if values.count(value) > 1})
            if duplicates:
                errors.append(
                    f"{name}: alternate key ({', '.join(key_columns)}) has duplicate values: "
                    f"{duplicates}."
                )

    for name, rows in rows_by_table.items():
        table = tables[name]
        for lookup in table.get("lookups", []):
            target_keys = keys_by_table.get(lookup["target_table"], set())
            for row_number, row in enumerate(rows, start=2):
                value = row.get(lookup["column"])
                if not value:
                    if lookup.get("required"):
                        errors.append(
                            f"{name}!{row_number}: required lookup {lookup['column']} is blank."
                        )
                    continue
                if str(value) not in target_keys and str(value) not in (lookup.get("existing_ids") or []):
                    errors.append(
                        f"{name}!{row_number}: lookup {lookup['column']} points to "
                        f"unknown {lookup['target_table']} key {value!r}."
                    )
    return rows_by_table, errors


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate an edited test-data workbook offline.")
    parser.add_argument("--workbook", required=True)
    args = parser.parse_args()
    emit_progress(
        6,
        "running",
        "Checking required fields, meaningful-data profiles, types, choices, keys, and lookups "
        "without writing to Dataverse.",
        "Show the approval summary if validation succeeds.",
    )
    rows_by_table, errors = validate(Path(args.workbook))
    if errors:
        emit_progress(
            6,
            "blocked",
            f"Found {len(errors)} workbook validation error(s); approval and writes remain blocked.",
            "Correct the listed workbook values and validate it again.",
        )
        print("Workbook is invalid:")
        for error in errors:
            print(f"- {error}")
        raise SystemExit(1)
    counts = ", ".join(f"{name}={len(rows)}" for name, rows in rows_by_table.items())
    emit_progress(
        6,
        "complete",
        f"The workbook is valid ({counts}). No Dataverse records have been written.",
        "Review the summary and explicitly approve or adjust the preview.",
    )
    print(f"Workbook is valid: {counts}")


if __name__ == "__main__":
    main()
